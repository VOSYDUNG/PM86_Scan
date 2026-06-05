from pathlib import Path
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "MASTER_GUIDE_PM86.md"
OUT_DOCX = ROOT / "docs" / "MASTER_GUIDE_PM86.docx"
OUT_DOCS = ROOT / "docs" / "MASTER_GUIDE_PM86.docs"
MIN_VALID_IMAGE_SIZE = 5000


def is_table_separator(line: str) -> bool:
    text = line.strip()
    if not text.startswith("|"):
        return False
    text = text.strip("|").strip()
    parts = [p.strip() for p in text.split("|")]
    if not parts:
        return False
    return all(set(p) <= set("-: ") and "-" in p for p in parts)


def split_table_row(line: str):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def apply_doc_style(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)


def build_doc(md_text: str) -> Document:
    doc = Document()
    apply_doc_style(doc)

    lines = md_text.splitlines()
    image_re = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
    heading_re = re.compile(r"^(#{1,6})\s+(.*)$")
    ordered_re = re.compile(r"^\d+\.\s+(.*)$")

    in_code = False
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue

        if in_code:
            p = doc.add_paragraph(line)
            p.style = doc.styles["No Spacing"]
            i += 1
            continue

        hm = heading_re.match(line)
        if hm:
            level = min(len(hm.group(1)), 4)
            doc.add_heading(hm.group(2).strip(), level=level)
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            header = split_table_row(line)
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(split_table_row(lines[i]))
                i += 1

            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for c, v in enumerate(header):
                table.cell(0, c).text = v
            for row in body:
                cells = table.add_row().cells
                for c in range(len(header)):
                    cells[c].text = row[c] if c < len(row) else ""
            continue

        im = image_re.match(line.strip())
        if im:
            alt, rel = im.groups()
            img = (MD_PATH.parent / rel).resolve()
            if img.exists() and img.stat().st_size >= MIN_VALID_IMAGE_SIZE:
                doc.add_paragraph(alt if alt else img.name)
                try:
                    doc.add_picture(str(img), width=Inches(5.8))
                except Exception:
                    doc.add_paragraph(f"[Khong chen duoc anh: {img.name}]")
            else:
                doc.add_paragraph(f"[Anh chua cap nhat: {alt or rel}]")
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(line[2:].strip())
            p.style = "List Bullet"
            i += 1
            continue

        om = ordered_re.match(line)
        if om:
            p = doc.add_paragraph(om.group(1).strip())
            p.style = "List Number"
            i += 1
            continue

        if not line.strip():
            doc.add_paragraph("")
            i += 1
            continue

        doc.add_paragraph(line)
        i += 1

    return doc


def main():
    text = MD_PATH.read_text(encoding="utf-8")
    doc = build_doc(text)
    candidates = [
        ROOT / "docs" / "MASTER_GUIDE_PM86_v2.docx",
        ROOT / "docs" / "MASTER_GUIDE_PM86.docx",
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    candidates.append(ROOT / "docs" / f"MASTER_GUIDE_PM86_{ts}.docx")

    last_err = None
    for out in candidates:
        try:
            doc.save(out)
            print(f"CREATED: {out}")
            return
        except PermissionError as e:
            last_err = e
            continue

    if last_err:
        raise last_err


if __name__ == "__main__":
    main()
